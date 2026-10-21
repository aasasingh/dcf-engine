"""The Excel formulas and the Python model are separate implementations; they must agree."""
import re

import pytest
from openpyxl import load_workbook
from pycel import ExcelCompiler

from dcf.assumptions import scenarios
from dcf.excel import write_workbook
from dcf.fetch import MarketData
from dcf.model import run_dcf
from dcf.pipeline import Analysis
from dcf.reverse import reverse_dcf
from dcf.sensitivity import scenario_table, tornado, wacc_terminal_grid


def _inline_defined_names(src, dst):
    """pycel can't read openpyxl 3.1 defined names, so substitute them with cell refs."""
    wb = load_workbook(src)
    names = {n: d.attr_text for n, d in wb.defined_names.items()}
    pattern = re.compile(r"\b(" + "|".join(sorted(names, key=len, reverse=True)) + r")\b")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.value = pattern.sub(lambda m: names[m.group(1)], c.value)
    for n in list(names):
        del wb.defined_names[n]
    wb.save(dst)


def _analysis(a, base, hist):
    md = MarketData(price=base.price, risk_free=a.risk_free, beta_raw=a.beta,
                    beta_adjusted=a.beta, beta_observations=60, as_of="2026-01-01")
    latest = hist[2025].copy()
    latest.attrs["label"] = "FY2025"
    return Analysis("TOY", "Toy Co", hist, latest, md, base, a, None, run_dcf(a, base),
                    reverse_dcf(a, base), wacc_terminal_grid(a, base), tornado(a, base),
                    scenario_table(scenarios(a), base))


@pytest.mark.parametrize("changes", [
    {},
    {"wacc_override": None, "mid_year_convention": True, "forecast_years": 10,
     "margin_start": 0.15, "margin_target": 0.25, "margin_fade_years": 5},
    {"wacc_override": None, "leases_as_debt": True, "forecast_years": 7},
    {"wacc_override": None, "forecast_years": 6, "growth_year1": -0.05},
])
def test_workbook_value_matches_python(tmp_path, toy_assumptions, toy_base, toy_hist, changes):
    a = toy_assumptions.with_(**changes)
    base = toy_base.__class__(**{**toy_base.__dict__, "debt": 20.0})
    x = _analysis(a, base, toy_hist)

    path = write_workbook(x, tmp_path / "model.xlsx")
    flat = tmp_path / "flat.xlsx"
    _inline_defined_names(path, flat)

    xl = ExcelCompiler(filename=str(flat))
    value_cell = next(f"DCF!B{r}" for r in range(20, 35)
                      if load_workbook(flat)["DCF"][f"A{r}"].value == "Value per share")
    assert xl.evaluate(value_cell) == pytest.approx(x.result.value_per_share, rel=1e-9)
